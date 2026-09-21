"""
excel_builder.py
生成标准差旅报销审批单
功能特性：
1. A列“人员”多行自动纵向合并居中
2. B列“序号”全局连续向下顺延 (1, 2, 3, 4, 5...)
3. 每个人员名下有独立的【合计】小计行 (SUM公式)
4. 最底端输出全局【汇总】大计行与跨人求和公式
5. 冻结表头以便于长流水滚动浏览
"""

from collections import defaultdict
from pathlib import Path
import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


def build_expense_excel(
    tickets: list[dict],
    output_path: str,
    project_name: str = "差旅报销",
    daily_meal_allowance: float = 0.0,
):
    """根据解析结果生成左侧合并人员、顺延排号的多人矩阵报销单。"""
    valid_tickets = [t for t in tickets if t.get("valid") and t.get("date")]
    if not valid_tickets:
        print("[!] 没有可用的票据数据，跳过 Excel 生成。")
        return

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "差旅报销审批单"
    ws.views.sheetView[0].showGridLines = True

    # 样式配置
    font_title = Font(name="微软雅黑", size=14, bold=True, color="000000")
    font_header = Font(name="微软雅黑", size=10, bold=True, color="FFFFFF")
    font_body = Font(name="微软雅黑", size=10)
    font_subtotal = Font(name="微软雅黑", size=10, bold=True)
    font_total = Font(name="微软雅黑", size=11, bold=True)

    fill_header = PatternFill(
        start_color="2C3E50", end_color="2C3E50", fill_type="solid"
    )
    fill_subtotal = PatternFill(
        start_color="F8F9F9", end_color="F8F9F9", fill_type="solid"
    )
    fill_total = PatternFill(
        start_color="EAEDED", end_color="EAEDED", fill_type="solid"
    )

    thin_border = Border(
        left=Side(style="thin", color="BDC3C7"),
        right=Side(style="thin", color="BDC3C7"),
        top=Side(style="thin", color="BDC3C7"),
        bottom=Side(style="thin", color="BDC3C7"),
    )

    # 财务数值标准格式：为 0 时安全展示为 "-"
    money_format = '_(* #,##0.00_);_(* (#,##0.00);_(* "-"??_);_(@_)'

    # 1. 标题行 (第1行，A1:I1 合并)
    ws.merge_cells("A1:I1")
    title_text = (
        f"{project_name.strip()} - 差旅报销审批单"
        if project_name.strip()
        else "差旅报销审批单"
    )
    title_cell = ws.cell(row=1, column=1, value=title_text)
    title_cell.font = font_title
    title_cell.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 36

    # 2. 表头行 (第2行，共9列)
    headers = [
        "人员",
        "序号",
        "日期",
        "打车",
        "高铁",
        "住宿",
        "餐补",
        "当日小计",
        "行程与凭证说明",
    ]
    ws.append(headers)
    ws.row_dimensions[2].height = 26

    for col_idx in range(1, 10):
        c = ws.cell(row=2, column=col_idx)
        c.font = font_header
        c.fill = fill_header
        c.alignment = Alignment(horizontal="center", vertical="center")
        c.border = thin_border

    # 3. 按 人员 -> 日期 聚类数据
    person_grouped = defaultdict(
        lambda: defaultdict(
            lambda: {
                "打车": 0.0,
                "高铁": 0.0,
                "住宿": 0.0,
                "descs": [],
            }
        )
    )

    for t in valid_tickets:
        p_name = str(t.get("person", "报销人")).strip() or "报销人"
        t_date = str(t.get("date", "")).replace("-", ".").strip()
        if not t_date:
            continue
        c_type = t.get("type", "其它")
        c_amt = float(t.get("amount", 0.0))
        c_desc = str(t.get("desc", "")).strip()

        if c_type in ["打车", "高铁", "住宿"]:
            person_grouped[p_name][t_date][c_type] += c_amt
        if c_desc and c_desc not in person_grouped[p_name][t_date]["descs"]:
            person_grouped[p_name][t_date]["descs"].append(c_desc)

    curr_row = 3
    seq_no = 1  # 全局顺延序号
    subtotal_rows = []  # 保存各人员合计行号，方便最后大汇总计算

    # 4. 逐个写入每个人员的数据块
    for p_name, dates_dict in person_grouped.items():
        sorted_dates = sorted(dates_dict.keys())
        person_start_row = curr_row

        for d_str in sorted_dates:
            data = dates_dict[d_str]
            ws.row_dimensions[curr_row].height = 22

            # B列: 序号
            c_seq = ws.cell(row=curr_row, column=2, value=seq_no)
            c_seq.alignment = Alignment(horizontal="center", vertical="center")
            seq_no += 1

            # C列: 日期
            c_date = ws.cell(row=curr_row, column=3, value=d_str)
            c_date.alignment = Alignment(horizontal="center", vertical="center")

            # D/E/F/G列: 打车、高铁、住宿、餐补
            taxi_v = data["打车"]
            train_v = data["高铁"]
            hotel_v = data["住宿"]
            meal_v = daily_meal_allowance if daily_meal_allowance > 0 else 0.0

            for col_idx, val in enumerate([taxi_v, train_v, hotel_v, meal_v], start=4):
                c = ws.cell(row=curr_row, column=col_idx)
                c.value = val
                c.number_format = money_format
                c.alignment = Alignment(horizontal="right", vertical="center")

            # H列: 当日小计公式 (=SUM(D{r}:G{r}))
            c_sub = ws.cell(row=curr_row, column=8, value=f"=SUM(D{curr_row}:G{curr_row})")
            c_sub.number_format = money_format
            c_sub.alignment = Alignment(horizontal="right", vertical="center")

            # I列: 行程与凭证说明
            desc_str = "; ".join(data["descs"]) if data["descs"] else "-"
            c_desc = ws.cell(row=curr_row, column=9, value=desc_str)
            c_desc.alignment = Alignment(horizontal="left", vertical="center")

            for col_idx in range(1, 10):
                cell = ws.cell(row=curr_row, column=col_idx)
                cell.border = thin_border
                cell.font = font_body

            curr_row += 1

        person_end_row = curr_row - 1

        # A列: 跨行合并人员名称
        if person_start_row < person_end_row:
            ws.merge_cells(
                start_row=person_start_row,
                start_column=1,
                end_row=person_end_row,
                end_column=1,
            )
        c_person = ws.cell(row=person_start_row, column=1, value=p_name)
        c_person.alignment = Alignment(horizontal="center", vertical="center")
        c_person.font = font_body

        # 写入个人【合计】小计行
        subtotal_row = curr_row
        subtotal_rows.append(subtotal_row)
        ws.row_dimensions[subtotal_row].height = 23

        ws.cell(row=subtotal_row, column=2, value="合计").alignment = Alignment(
            horizontal="center", vertical="center"
        )
        ws.cell(
            row=subtotal_row, column=3, value=f"共 {len(sorted_dates)} 天"
        ).alignment = Alignment(horizontal="center", vertical="center")

        # 对打车(D)、高铁(E)、住宿(F)、餐补(G)、当日小计(H)求和
        for col_idx, col_letter in enumerate(["D", "E", "F", "G", "H"], start=4):
            c = ws.cell(row=subtotal_row, column=col_idx)
            c.value = f"=SUM({col_letter}{person_start_row}:{col_letter}{person_end_row})"
            c.number_format = money_format
            c.alignment = Alignment(horizontal="right", vertical="center")

        ws.cell(row=subtotal_row, column=9, value="汇总").alignment = Alignment(
            horizontal="left", vertical="center"
        )

        for col_idx in range(1, 10):
            c = ws.cell(row=subtotal_row, column=col_idx)
            c.border = thin_border
            c.font = font_subtotal
            c.fill = fill_subtotal

        # 人员模块之间空出 1 行间隔
        curr_row += 2

    # 5. 全局项目【汇总】行 (底端大计)
    total_row = curr_row - 1
    ws.row_dimensions[total_row].height = 26

    ws.cell(row=total_row, column=1, value="汇总").alignment = Alignment(
        horizontal="center", vertical="center"
    )

    for col_idx, col_letter in enumerate(["D", "E", "F", "G", "H"], start=4):
        c = ws.cell(row=total_row, column=col_idx)
        if subtotal_rows:
            sub_cells = ",".join([f"{col_letter}{r}" for r in subtotal_rows])
            c.value = f"=SUM({sub_cells})"
        else:
            c.value = 0.0
        c.number_format = money_format
        c.alignment = Alignment(horizontal="right", vertical="center")

    for col_idx in range(1, 10):
        c = ws.cell(row=total_row, column=col_idx)
        c.border = thin_border
        c.font = font_total
        c.fill = fill_total

    # 6. 统一设置列宽
    column_widths = {
        "A": 10,
        "B": 8,
        "C": 14,
        "D": 12,
        "E": 12,
        "F": 12,
        "G": 12,
        "H": 14,
        "I": 42,
    }
    for col_letter, width in column_widths.items():
        ws.column_dimensions[col_letter].width = width

    # 7. 冻结表头
    ws.freeze_panes = "A3"

    Path(output_path).parent.mkdir(parents=True, exist_ok=True)
    wb.save(output_path)
    print(f"[✓] 矩阵报销表已生成: {output_path}")